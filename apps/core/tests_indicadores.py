"""Indicadores e apresentação: página, filtro de período, .pptx e base vazia."""
from io import BytesIO

from django.core.management import call_command
from django.test import TestCase
from pptx import Presentation

from apps.accounts.models import TipoUsuario, User
from apps.core.apresentacao import gerar_apresentacao
from apps.core.indicadores import Periodo, calcular, ler_periodo


class IndicadoresTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("seed_demo", verbosity=0)

    def setUp(self):
        self.client.login(username="analista", password="sst360demo")

    def test_pagina_abre_com_graficos(self):
        r = self.client.get("/indicadores/")
        self.assertEqual(r.status_code, 200)
        html = r.content.decode()
        self.assertIn('class="grafico"', html)
        self.assertIn("Baixar apresentação", html)

    def test_periodo_invertido_e_invalido(self):
        p = ler_periodo({"de": "2026-10-01", "ate": "2026-01-01"})
        self.assertLess(p.de, p.ate)
        p2 = ler_periodo({"de": "xx"})
        self.assertEqual(p2.de.day, 1)

    def test_apresentacao_baixa_e_abre(self):
        r = self.client.get("/indicadores/apresentacao.pptx?de=2026-05-01&ate=2026-10-15")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r["Content-Disposition"])
        prs = Presentation(BytesIO(r.content))
        self.assertGreaterEqual(len(prs.slides), 10)
        titulos = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.name == "Título"]
        self.assertIn("Resumo executivo", titulos)
        self.assertIn("FAP dos clientes", titulos)
        graficos = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
        self.assertGreaterEqual(graficos, 8)

    def test_periodo_sem_vendas(self):
        d = calcular(Periodo(*(ler_periodo({"de": "2019-01-01", "ate": "2019-02-28"}).__dict__.values())))
        self.assertEqual(d["comercial"]["vendas"]["total"], 0)
        prs = Presentation(gerar_apresentacao(d, "teste"))
        self.assertGreaterEqual(len(prs.slides), 10)
        self.assertEqual(self.client.get("/indicadores/?de=2019-01-01&ate=2019-02-28").status_code, 200)

    def test_cliente_nao_acessa(self):
        self.client.logout()
        self.client.login(username="cliente", password="sst360demo")
        self.assertEqual(self.client.get("/indicadores/").status_code, 403)
        self.assertEqual(self.client.get("/indicadores/apresentacao.pptx").status_code, 403)


class BaseVaziaTest(TestCase):
    def test_sem_dados(self):
        User.objects.create_user("equipe", password="x", tipo=TipoUsuario.INTERNO)
        self.client.login(username="equipe", password="x")
        self.assertEqual(self.client.get("/indicadores/").status_code, 200)
        r = self.client.get("/indicadores/apresentacao.pptx")
        self.assertEqual(r.status_code, 200)
        Presentation(BytesIO(r.content))
