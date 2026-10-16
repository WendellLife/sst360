"""Apresentação dos indicadores em PowerPoint (.pptx), com gráficos nativos.

Os gráficos são do próprio PowerPoint (dá para editar dados e cores depois) e cada
slide leva uma mensagem no título. A fonte é Arial para o arquivo abrir igual em
qualquer computador; as cores são as da identidade Razach.
"""
from io import BytesIO
from pathlib import Path

from django.conf import settings
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .indicadores import moeda

# ---------------------------------------------------------------- identidade
GRAFITE = RGBColor(0x0C, 0x17, 0x19)
PETROLEO = RGBColor(0x1B, 0x64, 0x70)
TURQUESA = RGBColor(0x35, 0x98, 0x9F)
CIANO = RGBColor(0x6F, 0xCB, 0xCE)
NEVOA = RGBColor(0xF4, 0xF7, 0xF7)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
TINTA = RGBColor(0x1F, 0x2E, 0x30)
SUAVE = RGBColor(0x4A, 0x5A, 0x5C)
FRACA = RGBColor(0x66, 0x78, 0x7A)
LINHA = RGBColor(0xD5, 0xDE, 0xDE)
VERDE = RGBColor(0x1B, 0x7F, 0x63)
AMARELO = RGBColor(0xA9, 0x74, 0x1A)
VERMELHO = RGBColor(0xB0, 0x3A, 0x32)
CINZA = RGBColor(0xB6, 0xC3, 0xC4)
COR_NIVEL = {"BMAX": VERDE, "BON": VERDE, "NEU": RGBColor(0x66, 0x78, 0x7A), "MAL": AMARELO, "MMAX": VERMELHO}
FONTE = "Arial"
L, A = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.6)
IMG = Path(settings.BASE_DIR) / "static" / "razach" / "img"


# ---------------------------------------------------------------- utilitários
def _texto(slide, x, y, w, h, texto, *, tam=14, negrito=False, cor=TINTA, alinhar=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP, nome=None):
    caixa = slide.shapes.add_textbox(x, y, w, h)
    if nome:
        caixa.name = nome
    tf = caixa.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancora
    linhas = texto if isinstance(texto, list) else [texto]
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhar
        partes = linha if isinstance(linha, list) else [(linha, {})]
        for conteudo, estilo in partes:
            r = p.add_run()
            r.text = conteudo
            f = r.font
            f.name = FONTE
            f.size = Pt(estilo.get("tam", tam))
            f.bold = estilo.get("negrito", negrito)
            f.color.rgb = estilo.get("cor", cor)
        if i:
            p.space_before = Pt(6)
    return caixa


def _caixa(slide, x, y, w, h, cor=NEVOA, nome=None):
    forma = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    forma.adjustments[0] = 0.06
    forma.fill.solid()
    forma.fill.fore_color.rgb = cor
    forma.line.fill.background()
    forma.shadow.inherit = False
    if nome:
        forma.name = nome
    return forma


def _fundo(slide, cor):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = cor


def _numero(slide, x, y, w, rotulo, valor, sub="", *, cor_valor=GRAFITE, tam=30, fundo=NEVOA):
    _caixa(slide, x, y, w, Inches(1.45), fundo)
    pad = Inches(0.22)
    _texto(slide, x + pad, y + Inches(0.18), w - 2 * pad, Inches(0.3), rotulo.upper(), tam=10, negrito=True, cor=FRACA)
    _texto(slide, x + pad, y + Inches(0.46), w - 2 * pad, Inches(0.55), valor, tam=tam, negrito=True, cor=cor_valor)
    if sub:
        _texto(slide, x + pad, y + Inches(1.04), w - 2 * pad, Inches(0.3), sub, tam=11, cor=SUAVE)


class Deck:
    def __init__(self, dados, autor):
        self.d = dados
        self.autor = autor
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = L, A
        self.branco = self.prs.slide_layouts[6]
        self.n = 0

    def novo(self, titulo, mensagem="", *, escuro=False, notas=""):
        slide = self.prs.slides.add_slide(self.branco)
        self.n += 1
        _fundo(slide, GRAFITE if escuro else BRANCO)
        if titulo:
            _texto(slide, MARGEM, Inches(0.45), L - 2 * MARGEM, Inches(0.6), titulo, tam=30, negrito=True,
                   cor=BRANCO if escuro else GRAFITE, nome="Título")
        if mensagem:
            _texto(slide, MARGEM, Inches(1.1), L - 2 * MARGEM, Inches(0.45), mensagem, tam=15,
                   cor=CIANO if escuro else SUAVE, nome="Mensagem")
        if not escuro:
            _texto(slide, MARGEM, A - Inches(0.45), Inches(9), Inches(0.25),
                   f"SST360 · Indicadores da carteira · {self.d['periodo'].rotulo}", tam=9, cor=FRACA)
            _texto(slide, L - MARGEM - Inches(1), A - Inches(0.45), Inches(1), Inches(0.25), str(self.n), tam=9,
                   cor=FRACA, alinhar=PP_ALIGN.RIGHT)
        if notas:
            slide.notes_slide.notes_text_frame.text = notas
        return slide

    # ------------------------------------------------------------ gráficos nativos
    def _estilo_grafico(self, grafico, *, formato="#,##0", legenda=False, rotulos=True, eixo_valor=True, tam=11):
        grafico.has_title = False
        grafico.font.name = FONTE
        grafico.font.size = Pt(tam)
        grafico.font.color.rgb = SUAVE
        grafico.has_legend = legenda
        if legenda:
            grafico.legend.position = XL_LEGEND_POSITION.BOTTOM
            grafico.legend.include_in_layout = False
            grafico.legend.font.size = Pt(tam)
        plot = grafico.plots[0]
        if rotulos:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = "0;-0;;" if formato == "0" else formato  # contagens: esconde os zeros
            dl.number_format_is_linked = False
            dl.font.size = Pt(tam)
            dl.font.bold = True
            dl.font.color.rgb = TINTA
        try:
            va = grafico.value_axis
            va.has_major_gridlines = eixo_valor
            if eixo_valor:
                va.major_gridlines.format.line.color.rgb = LINHA
                va.major_gridlines.format.line.width = Pt(0.75)
            va.format.line.fill.background()
            va.tick_labels.number_format = formato
            va.tick_labels.number_format_is_linked = False
            va.visible = eixo_valor
            ca = grafico.category_axis
            ca.format.line.color.rgb = LINHA
            ca.has_major_gridlines = False
        except (ValueError, AttributeError):
            pass
        return grafico

    def colunas(self, slide, x, y, w, h, categorias, valores, *, formato='"R$" #,##0', destaque_ultimo=True, cor=CIANO, nome_serie="Total"):
        dados = CategoryChartData()
        dados.categories = categorias
        dados.add_series(nome_serie, [float(v) for v in valores])
        g = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, dados).chart
        self._estilo_grafico(g, formato=formato)
        plot = g.plots[0]
        plot.gap_width = 120
        serie = plot.series[0]
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = cor
        if destaque_ultimo and valores:
            ponto = serie.points[len(valores) - 1]
            ponto.format.fill.solid()
            ponto.format.fill.fore_color.rgb = PETROLEO
        return g

    def barras(self, slide, x, y, w, h, categorias, valores, *, formato='"R$" #,##0', cores=None, cor=TURQUESA):
        dados = CategoryChartData()
        # barras horizontais desenham de baixo para cima: inverte para o maior ficar em cima
        dados.categories = list(reversed(categorias))
        dados.add_series("Total", [float(v) for v in reversed(valores)])
        g = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, x, y, w, h, dados).chart
        self._estilo_grafico(g, formato=formato, eixo_valor=False)
        plot = g.plots[0]
        plot.gap_width = 70
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        serie = plot.series[0]
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = cor
        if cores:
            for i, c in enumerate(reversed(cores)):
                p = serie.points[i]
                p.format.fill.solid()
                p.format.fill.fore_color.rgb = c
        return g

    def rosca(self, slide, x, y, w, h, categorias, valores, cores):
        pares = [(c, v, cor) for c, v, cor in zip(categorias, valores, cores) if v]
        if not pares:
            _texto(slide, x, y, w, Inches(0.4), "Sem dados.", tam=12, cor=FRACA)
            return None
        dados = CategoryChartData()
        dados.categories = [f"{p[0]} ({int(p[1])})" for p in pares]
        dados.add_series("Total", [float(p[1]) for p in pares])
        g = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y, w, h, dados).chart
        self._estilo_grafico(g, formato="0", legenda=True)
        plot = g.plots[0]
        plot.data_labels.font.color.rgb = BRANCO
        for i, (_, _, cor) in enumerate(pares):
            p = plot.series[0].points[i]
            p.format.fill.solid()
            p.format.fill.fore_color.rgb = cor
            p.format.line.color.rgb = BRANCO
            p.format.line.width = Pt(2)
        return g

    def tabela(self, slide, x, y, w, cabecalho, linhas, larguras, *, altura_linha=Inches(0.48), tam=None):
        tam = tam or (13 if len(linhas) <= 5 else 11)
        forma = slide.shapes.add_table(len(linhas) + 1, len(cabecalho), x, y, w, altura_linha * (len(linhas) + 1))
        t = forma.table
        total = sum(larguras)
        for i, lw in enumerate(larguras):
            t.columns[i].width = Emu(int(w * lw / total))
        for r in range(len(linhas) + 1):
            t.rows[r].height = altura_linha
            for c in range(len(cabecalho)):
                cel = t.cell(r, c)
                cel.fill.solid()
                cel.fill.fore_color.rgb = PETROLEO if r == 0 else (BRANCO if r % 2 else NEVOA)
                cel.margin_left = cel.margin_right = Inches(0.1)
                cel.vertical_anchor = MSO_ANCHOR.MIDDLE
                valor = cabecalho[c] if r == 0 else linhas[r - 1][c]
                cor = BRANCO if r == 0 else TINTA
                negrito = r == 0
                if isinstance(valor, tuple):
                    valor, cor = valor
                tf = cel.text_frame
                tf.word_wrap = True
                tf.paragraphs[0].text = ""
                run = tf.paragraphs[0].add_run()
                run.text = str(valor)
                run.font.name = FONTE
                run.font.size = Pt(tam if r else tam - 1)
                run.font.bold = negrito
                run.font.color.rgb = cor
        return forma


# ---------------------------------------------------------------- slides
def _capa(dk):
    d = dk.d
    s = dk.novo("", escuro=True, notas="Abertura: período, quem preparou e o que a apresentação cobre.")
    if (IMG / "razach-simbolo-claro.png").exists():
        s.shapes.add_picture(str(IMG / "razach-simbolo-claro.png"), MARGEM, Inches(0.7), height=Inches(0.7))
    _texto(s, MARGEM, Inches(2.2), Inches(10), Inches(0.4), "SST360 · UNIVERSO FX · LIFE LABORAL", tam=13, negrito=True, cor=CIANO)
    _texto(s, MARGEM, Inches(2.7), Inches(11), Inches(1.8), "Indicadores da carteira de saúde e segurança do trabalho", tam=40, negrito=True, cor=BRANCO)
    _texto(s, MARGEM, Inches(4.75), Inches(11), Inches(0.5), d["periodo"].rotulo, tam=20, cor=NEVOA)
    _texto(s, MARGEM, Inches(5.25), Inches(11), Inches(0.4),
           f"Comercial, conformidade, FAP e carteira · gerado em {d['gerado_em']:%d/%m/%Y %H:%M}" + (f" por {dk.autor}" if dk.autor else ""),
           tam=13, cor=CINZA)
    if (IMG / "razach-combinado-claro.png").exists():
        s.shapes.add_picture(str(IMG / "razach-combinado-claro.png"), L - MARGEM - Inches(1.9), A - Inches(1.0), height=Inches(0.52))


def _resumo(dk):
    d = dk.d
    v, c, f, cf = d["comercial"]["vendas"], d["carteira"], d["fap"], d["conformidade"]
    s = dk.novo("Resumo executivo", f"{moeda(v['total'])} vendidos no período e {moeda(f['potencial'])}/ano de economia possível no FAP dos clientes",
                notas="Os quatro números que resumem o período e as frases de destaque, separadas em atenção, oportunidade e resultado.")
    w = (L - 2 * MARGEM - Inches(0.3) * 3) / 4
    y = Inches(1.8)
    cards = [
        ("Vendido no período", moeda(v["total"]), f"{v['quantidade']} venda(s) · {moeda(v['recebido'])} recebido", PETROLEO),
        ("Receita recorrente", moeda(c["receita_mensal"]), f"por mês · {c['gestoes']} gestão(ões) ativa(s)", GRAFITE),
        ("Economia possível no FAP", moeda(f["potencial"]), f"por ano · {f['nao_contratado']} cliente(s) sem o serviço", VERDE),
        ("Documentos em dia", f"{cf['documentos']['pct_vigentes']}%", f"{cf['documentos']['vencido']} vencido(s) · plano {cf['acoes']['pct']}% concluído",
         VERMELHO if cf["documentos"]["vencido"] else GRAFITE),
    ]
    for i, (rot, val, sub, cor) in enumerate(cards):
        _numero(s, MARGEM + i * (w + Inches(0.3)), y, w, rot, val, sub, cor_valor=cor, tam=24)
    y2 = Inches(3.65)
    _texto(s, MARGEM, y2, Inches(6), Inches(0.35), "DESTAQUES DO PERÍODO", tam=11, negrito=True, cor=FRACA)
    grupos = [("alerta", "Atenção", VERMELHO), ("oportunidade", "Oportunidade", AMARELO), ("resultado", "Resultado", VERDE)]
    col_w = (L - 2 * MARGEM - Inches(0.6)) / 3
    for k, (tipo, rotulo, cor) in enumerate(grupos):
        x = MARGEM + k * (col_w + Inches(0.3))
        frases = [frase for t, frase in d["destaques"] if t == tipo] or ["Nada a destacar no período."]
        _caixa(s, x, y2 + Inches(0.45), col_w, Inches(2.75), NEVOA)
        marca = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.25), y2 + Inches(0.72), Inches(0.16), Inches(0.16))
        marca.fill.solid()
        marca.fill.fore_color.rgb = cor
        marca.line.fill.background()
        _texto(s, x + Inches(0.5), y2 + Inches(0.64), col_w - Inches(0.7), Inches(0.35), rotulo, tam=14, negrito=True, cor=GRAFITE)
        _texto(s, x + Inches(0.25), y2 + Inches(1.1), col_w - Inches(0.5), Inches(2.0), frases, tam=13, cor=TINTA)


def _vendas(dk):
    v = dk.d["comercial"]["vendas"]
    meses = v["por_mes"]
    melhor = max(meses, key=lambda m: m["total"]) if meses else None
    msg = f"{moeda(v['total'])} em {v['quantidade']} venda(s)" + (f"; melhor mês: {melhor['rotulo']} com {moeda(melhor['total'])}" if melhor and melhor["total"] else "")
    s = dk.novo("Vendas por mês", msg, notas="Colunas: valor vendido por mês, sem as canceladas. O mês mais recente vem destacado.")
    dk.colunas(s, MARGEM, Inches(1.75), Inches(8.3), Inches(5.0), [m["rotulo"] for m in meses], [m["total"] for m in meses])
    x = MARGEM + Inches(8.7)
    w = L - MARGEM - x
    itens = [("Recebido", moeda(v["recebido"]), "status recebida", VERDE), ("A receber", moeda(v["a_receber"]), "fechadas e faturadas", GRAFITE),
             ("Recorrente", moeda(v["recorrente"]), "vendas mensais", PETROLEO)]
    for i, (rot, val, sub, cor) in enumerate(itens):
        _numero(s, x, Inches(1.8) + i * Inches(1.65), w, rot, val, sub, cor_valor=cor, tam=22)


def _servicos(dk):
    v = dk.d["comercial"]["vendas"]
    itens = v["por_servico"]
    msg = (f"{itens[0]['rotulo']} lidera com {moeda(itens[0]['total'])}" if itens else "Sem vendas no período")
    s = dk.novo("Vendas por serviço", msg, notas="Barras: quanto cada linha do portfólio vendeu no período.")
    if itens:
        dk.barras(s, MARGEM, Inches(1.7), Inches(7.6), Inches(5.1), [i["rotulo"] for i in itens], [i["total"] for i in itens])
    x = MARGEM + Inches(8.0)
    w = L - MARGEM - x
    _texto(s, x, Inches(1.85), w, Inches(0.35), "MAIORES CLIENTES NO PERÍODO", tam=11, negrito=True, cor=FRACA)
    clientes = v["por_cliente"][:5]
    if clientes:
        dk.tabela(s, x, Inches(2.3), w, ["Cliente", "Vendido"], [[c["rotulo"], moeda(c["total"])] for c in clientes], [3, 2])
    _numero(s, x, Inches(5.15), w, "Ticket médio", moeda(v["ticket_medio"]), f"{v['canceladas']} venda(s) cancelada(s)", tam=22)


def _funil(dk):
    fu, dg = dk.d["comercial"]["funil"], dk.d["comercial"]["diagnosticos"]
    s = dk.novo("Funil de prospecção e diagnósticos",
                f"{fu['abertos']} lead(s) em aberto somando {moeda(fu['pipeline'])}/mês; conversão de {fu['conversao']}% no período",
                notas="Leads por estágio do funil e o resultado dos diagnósticos gratuitos feitos no site.")
    dk.barras(s, MARGEM, Inches(1.65), Inches(6.6), Inches(3.6), [e["rotulo"] for e in fu["por_estagio"]], [e["total"] for e in fu["por_estagio"]], formato="0")
    _numero(s, MARGEM, Inches(5.45), Inches(3.15), "Ganhos no período", str(fu["ganhos"]), f"{fu['perdidos']} perdido(s)", cor_valor=VERDE, tam=24)
    _numero(s, MARGEM + Inches(3.45), Inches(5.45), Inches(3.15), "Conversão", f"{fu['conversao']}%", "ganhos ÷ fechados", tam=24)
    x = MARGEM + Inches(7.1)
    w = L - MARGEM - x
    _texto(s, x, Inches(1.8), w, Inches(0.35), "DIAGNÓSTICOS DO SITE", tam=11, negrito=True, cor=FRACA)
    cores = {"ALTA": VERMELHO, "MEDIA": AMARELO, "BAIXA": VERDE}
    dk.barras(s, x, Inches(2.15), w, Inches(2.3), [u["rotulo"] for u in dg["por_urgencia"]], [u["total"] for u in dg["por_urgencia"]],
              formato="0", cores=[cores[u["codigo"]] for u in dg["por_urgencia"]])
    meio = (w - Inches(0.3)) / 2
    _numero(s, x, Inches(4.6), meio, "Diagnósticos", str(dg["total"]), f"{dg['pct_lead']}% viraram lead", tam=24)
    _numero(s, x + meio + Inches(0.3), Inches(4.6), meio, "Exposição média", moeda(dg["exposicao_media"]), "teto em multas", cor_valor=VERMELHO, tam=18)


def _conformidade(dk):
    cf = dk.d["conformidade"]
    docs, tre = cf["documentos"], cf["treinamentos"]
    s = dk.novo("Documentos e reciclagens",
                f"{docs['pct_vigentes']}% dos documentos técnicos em dia; {docs['vencido']} vencido(s) e {tre['vencidas']} reciclagem(ns) vencida(s)",
                notas="Situação de hoje das gestões ativas. Vigente inclui os que vencem em mais de 60 dias.")
    w = (L - 2 * MARGEM - Inches(0.4)) / 2
    for k, (titulo, cats, vals, cores, sub) in enumerate([
        ("Documentos técnicos", ["Vigentes", "A revisar em 60 dias", "Vencidos", "Em elaboração"],
         [docs["vigente"], docs["a_vencer"], docs["vencido"], docs["elaboracao"]], [VERDE, AMARELO, VERMELHO, CINZA], f"{docs['total']} documento(s)"),
        ("Reciclagens de treinamento", ["Em dia", "A reciclar em 60 dias", "Vencidas"],
         [tre["em_dia"], tre["a_vencer"], tre["vencidas"]], [VERDE, AMARELO, VERMELHO],
         f"{tre['pessoas_periodo']} pessoa(s) treinada(s) no período · {tre['horas_periodo']} h·pessoa"),
    ]):
        x = MARGEM + k * (w + Inches(0.4))
        _caixa(s, x, Inches(1.75), w, Inches(5.05), NEVOA)
        _texto(s, x + Inches(0.3), Inches(1.95), w - Inches(0.6), Inches(0.35), titulo, tam=16, negrito=True, cor=GRAFITE)
        _texto(s, x + Inches(0.3), Inches(2.35), w - Inches(0.6), Inches(0.3), sub, tam=11, cor=SUAVE)
        dk.rosca(s, x + Inches(0.3), Inches(2.75), w - Inches(0.6), Inches(3.9), cats, vals, cores)


def _prazos(dk):
    cf = dk.d["conformidade"]
    pz, ac = cf["prazos"], cf["acoes"]
    s = dk.novo("Prazos normativos e planos de ação",
                f"{pz['pct_no_prazo']}% dos prazos cumpridos no prazo; {pz['atrasados']} atrasado(s) e {ac['atrasadas']} ação(ões) vencida(s) hoje",
                notas="Prazos: vencimento dentro do período. Atrasados e próximos 30 dias são a situação de hoje.")
    w = (L - 2 * MARGEM - Inches(0.3) * 3) / 4
    itens = [("Prazos no período", str(pz["no_periodo"]), f"{pz['cumpridos']} cumprido(s)", GRAFITE),
             ("Cumpridos no prazo", f"{pz['pct_no_prazo']}%", f"{pz['no_prazo']} de {pz['cumpridos']}", VERDE),
             ("Atrasados agora", str(pz["atrasados"]), "pendentes vencidos", VERMELHO if pz["atrasados"] else GRAFITE),
             ("Próximos 30 dias", str(pz["proximos_30"]), "para agendar", AMARELO if pz["proximos_30"] else GRAFITE)]
    for i, (r, v, sub, cor) in enumerate(itens):
        _numero(s, MARGEM + i * (w + Inches(0.3)), Inches(1.8), w, r, v, sub, cor_valor=cor)
    _texto(s, MARGEM, Inches(3.6), Inches(6), Inches(0.35), "PLANOS DE AÇÃO DAS GESTÕES ATIVAS", tam=11, negrito=True, cor=FRACA)
    dados = [("Concluídas", ac["encerradas"], VERDE), ("Em aberto no prazo", max(ac["abertas"] - ac["atrasadas"], 0), CINZA), ("Atrasadas", ac["atrasadas"], VERMELHO)]
    dk.barras(s, MARGEM, Inches(4.0), Inches(7.4), Inches(2.6), [x[0] for x in dados], [x[1] for x in dados], formato="0", cores=[x[2] for x in dados])
    _numero(s, MARGEM + Inches(8.0), Inches(4.1), L - 2 * MARGEM - Inches(8.0), "Concluído", f"{ac['pct']}%",
            f"{ac['encerradas']} de {ac['total']} ações" + (f" · {moeda(ac['custo_aberto'])} em aberto" if ac["custo_aberto"] else ""), cor_valor=PETROLEO, tam=36)


def _fap(dk):
    f = dk.d["fap"]
    linhas = f["linhas"]
    s = dk.novo("FAP dos clientes",
                f"Economia possível de {moeda(f['potencial'])}/ano; os clientes em malus pagam {moeda(f['acrescimo'])}/ano acima do FAP neutro",
                notas="FAP vigente de cada cliente. Abaixo de 1,0000 é bônus (desconto no RAT); acima, malus (acréscimo). A cor segue o nível.")
    if linhas:
        g = dk.barras(s, MARGEM, Inches(1.7), Inches(7.4), Inches(4.0), [str(x["empresa"]) for x in linhas], [float(x["c"].fap) for x in linhas],
                      formato="0.0000", cores=[COR_NIVEL.get(x["c"].nivel["codigo"], FRACA) for x in linhas])
        va = g.value_axis
        va.minimum_scale, va.maximum_scale = 0, 2
        _texto(s, MARGEM, Inches(5.75), Inches(7.4), Inches(0.6),
               "Verde: bônus · cinza: neutro · âmbar: malus · vermelho: malus alto. FAP entre 0,5000 e 2,0000.", tam=11, cor=FRACA)
    x = MARGEM + Inches(7.9)
    w = L - MARGEM - x
    _numero(s, x, Inches(1.8), w, "Economia possível", moeda(f["potencial"]), "por ano, somando a carteira", cor_valor=VERDE, tam=24)
    _numero(s, x, Inches(3.45), w, "Pago acima do neutro", moeda(f["acrescimo"]), "RAT extra por ano", cor_valor=VERMELHO, tam=24)
    _numero(s, x, Inches(5.1), w, "Sem redução do FAP", f"{f['nao_contratado']} cliente(s)", f"{f['bloqueados']} com bonificação bloqueada", tam=24)


def _oportunidades(dk):
    f = dk.d["fap"]
    ops = f["oportunidades"]
    if not ops:
        return
    s = dk.novo("Onde agir no FAP", f"{len(ops)} cliente(s) com economia identificada, do maior potencial para o menor",
                notas="Argumentos prontos para a conversa com cada cliente. Economia possível = maior valor entre chegar à meta, voltar ao neutro e recuperar a bonificação bloqueada.")
    linhas = []
    for x in ops:
        linhas.append([str(x["empresa"]), f"{float(x['c'].fap):.4f}".replace(".", ",") + f" · {x['c'].nivel['rotulo']}",
                       (moeda(x["potencial"]) + "/ano", VERDE), ("Sim" if x["contratado"] else "Não", TINTA if x["contratado"] else AMARELO), x["argumento"]])
    dk.tabela(s, MARGEM, Inches(1.8), L - 2 * MARGEM, ["Cliente", "FAP · nível", "Economia possível", "Redução contratada", "Argumento"],
              linhas, [2.2, 1.6, 1.7, 1.3, 5.5], altura_linha=Inches(1.05))


def _clientes(dk):
    cl = dk.d["clientes"]
    if not cl:
        return
    com_pend = sum(1 for c in cl if c["alerta"])
    s = dk.novo("Situação por cliente", f"{com_pend} de {len(cl)} cliente(s) com pendências; quem tem mais aparece primeiro",
                notas="Pendências = documentos vencidos + prazos atrasados + ações atrasadas + reciclagens vencidas.")
    linhas = []
    for c in cl[:10]:
        ind = c["ind"]
        pend = []
        if ind["documentos_vencidos"]:
            pend.append(f"{ind['documentos_vencidos']} doc.")
        if c["prazos_atrasados"]:
            pend.append(f"{c['prazos_atrasados']} prazo(s)")
        if ind["acoes_atrasadas"]:
            pend.append(f"{ind['acoes_atrasadas']} ação(ões)")
        if ind["treinamentos_vencidos"]:
            pend.append(f"{ind['treinamentos_vencidos']} reciclagem(ns)")
        fap = c["fap"]["c"] if c["fap"] else None
        linhas.append([c["nome"], f"{c['etapa']} · {c['etapa_rotulo']}", f"{c['pct_docs']}%", f"{ind['acoes_percentual']}%",
                       (", ".join(pend) if pend else "em dia", VERMELHO if pend else VERDE),
                       (f"{float(fap.fap):.4f}".replace(".", ",") + f" · {fap.nivel['rotulo']}" if fap else "—",
                        COR_NIVEL.get(fap.nivel["codigo"], TINTA) if fap else FRACA)])
    dk.tabela(s, MARGEM, Inches(1.8), L - 2 * MARGEM, ["Cliente", "Etapa do método", "Documentos em dia", "Plano de ação", "Pendências", "FAP"],
              linhas, [2.4, 2.6, 1.3, 1.2, 3.0, 1.6], altura_linha=Inches(0.62 if len(linhas) <= 5 else 0.46))


def _carteira(dk):
    c = dk.d["carteira"]
    etapas = c["por_etapa"]
    mais = max(etapas, key=lambda e: e["total"]) if etapas else None
    s = dk.novo("Carteira por etapa do método",
                f"{c['gestoes']} gestão(ões) ativa(s) cobrindo {c['trabalhadores']} trabalhadores" + (f"; mais clientes em {mais['nome'].lower()}" if mais and mais["total"] else ""),
                notas="Em que etapa dos 12 passos do método cada gestão ativa está.")
    dk.colunas(s, MARGEM, Inches(1.75), Inches(8.6), Inches(5.0), [e["curto"] for e in etapas], [e["total"] for e in etapas],
               formato="0", destaque_ultimo=False, cor=TURQUESA)
    x = MARGEM + Inches(9.0)
    w = L - MARGEM - x
    _numero(s, x, Inches(1.8), w, "Receita recorrente", moeda(c["receita_mensal"]), "por mês", cor_valor=PETROLEO, tam=22)
    _numero(s, x, Inches(3.45), w, "Clientes", str(c["clientes"]), f"{c['gestoes']} gestão(ões) ativa(s)", tam=26)
    legenda = [f"{e['curto']} · {e['nome']}" for e in etapas if e["total"]]
    _texto(s, x, Inches(5.15), w, Inches(1.6), legenda or ["Sem gestões ativas."], tam=11, cor=SUAVE)


def _proximos(dk):
    d = dk.d
    acoes = []
    cf, f = d["conformidade"], d["fap"]
    if cf["prazos"]["atrasados"]:
        acoes.append(f"Regularizar {cf['prazos']['atrasados']} prazo(s) normativo(s) atrasado(s)")
    if cf["documentos"]["vencido"]:
        acoes.append(f"Revisar {cf['documentos']['vencido']} documento(s) vencido(s) (PGR, PCMSO, LTCAT, laudos)")
    if cf["treinamentos"]["vencidas"]:
        acoes.append(f"Agendar {cf['treinamentos']['vencidas']} reciclagem(ns) vencida(s)")
    if cf["acoes"]["atrasadas"]:
        acoes.append(f"Destravar {cf['acoes']['atrasadas']} ação(ões) do plano com prazo vencido")
    if f["oportunidades"]:
        top = f["oportunidades"][0]
        acoes.append(f"Apresentar a redução do FAP a {top['empresa']} ({moeda(top['potencial'])}/ano)")
    if cf["prazos"]["proximos_30"]:
        acoes.append(f"Preparar {cf['prazos']['proximos_30']} prazo(s) que vencem nos próximos 30 dias")
    acoes = acoes[:6] or ["Manter o acompanhamento: nenhuma pendência crítica no período."]
    s = dk.novo("Próximos passos", "O que fazer a partir destes números", escuro=True,
                notas="Gerado a partir das pendências e oportunidades encontradas no período.")
    w = (L - 2 * MARGEM - Inches(0.4)) / 2
    for i, texto in enumerate(acoes):
        col, lin = i % 2, i // 2
        x = MARGEM + col * (w + Inches(0.4))
        y = Inches(2.0) + lin * Inches(1.5)
        _caixa(s, x, y, w, Inches(1.25), RGBColor(0x12, 0x1F, 0x21))
        _texto(s, x + Inches(0.3), y + Inches(0.22), Inches(0.8), Inches(0.8), f"{i + 1:02d}", tam=28, negrito=True, cor=CIANO)
        _texto(s, x + Inches(1.2), y, w - Inches(1.5), Inches(1.25), texto, tam=16, cor=BRANCO, ancora=MSO_ANCHOR.MIDDLE)
    if (IMG / "razach-combinado-claro.png").exists():
        s.shapes.add_picture(str(IMG / "razach-combinado-claro.png"), L - MARGEM - Inches(1.9), A - Inches(0.9), height=Inches(0.52))


def gerar_apresentacao(dados: dict, autor: str = "") -> BytesIO:
    dk = Deck(dados, autor)
    dk.prs.core_properties.title = f"Indicadores da carteira · {dados['periodo'].rotulo}"
    dk.prs.core_properties.author = autor or "SST360"
    for montar in (_capa, _resumo, _vendas, _servicos, _funil, _conformidade, _prazos, _fap, _oportunidades, _clientes, _carteira, _proximos):
        montar(dk)
    saida = BytesIO()
    dk.prs.save(saida)
    saida.seek(0)
    return saida
